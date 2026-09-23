#!/usr/bin/env python3
"""
Basket Command Centre 2 -- dashboard category tracker.

Fetches the live data.json from https://moopheus1.github.io/Basket-Command-Centre-2/data.json,
replicates the dashboard's own client-side JS logic (computeMetrics, computeCrossStatus,
computeReboundMA50, bull-flag filter, EOD breakout filter, computeScores/combinedSignal)
in Python, logs today's category membership to an xlsx tracker, and resolves prior log
rows' forward returns (5/10/20 trading sessions later) using the bars now available in
the freshly-fetched data.json.

This is a forward test: every day a ticker appears in a category, this logs its close
price that day as "entry", then measures what the close was N trading sessions later.
It does NOT model an actual trade (no stop/target/sizing) -- it measures raw forward
price drift after each dashboard flag, categorised by which panel flagged it.
"""
import json
import math
import statistics
import sys
import urllib.request
from datetime import datetime, timezone
import openpyxl
from openpyxl import Workbook

DATA_URL = "https://moopheus1.github.io/Basket-Command-Centre-2/data.json"
TRACKER_PATH = sys.argv[1] if len(sys.argv) > 1 else "Basket-Command-Centre-tracker.xlsx"
EXPORT_DIR = sys.argv[sys.argv.index("--export") + 1] if "--export" in sys.argv else None

HORIZONS = [5, 10, 20]  # trading sessions forward
UP_THRESH, DOWN_THRESH = 0.5, -0.5  # % return thresholds for UP/DOWN/FLAT status

CROSS_LOOKBACK = 20
CROSS_APPROACH_GAP = 2
CROSS_CONVERGE_BARS = 5

MIN_RISK_ATR_MULT = 0.5
MIN_RR_MA50 = 1.5
REBOUND_HORIZON_DAYS = 15

BULL_FLAG_MIN_PULLBACK, BULL_FLAG_MAX_PULLBACK = 3, 14
BULL_FLAG_MAX_VOL_RATIO = 0.85
BULL_FLAG_RSI_MIN, BULL_FLAG_RSI_MAX = 40, 60

BREAKOUT_MIN_DAY_PCT = 5
BREAKOUT_MIN_VOL_RATIO = 2

REL_STRENGTH_BOUND = 10


def fetch_json(url):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=30) as r:
        return json.loads(r.read().decode("utf-8"))


def sma(closes, n):
    if len(closes) < n:
        return None
    s = closes[-n:]
    return sum(s) / n


def rsi14(closes):
    if len(closes) < 15:
        return None
    gains = losses = 0.0
    for i in range(1, 15):
        ch = closes[-i] - closes[-i - 1]
        if ch > 0:
            gains += ch
        else:
            losses -= ch
    avg_gain, avg_loss = gains / 14, losses / 14
    if avg_loss == 0:
        return 100.0
    rs = avg_gain / avg_loss
    return 100 - (100 / (1 + rs))


def atr14(bars):
    # bars: list of dict(open,high,low,close,volume)
    if len(bars) < 15:
        return None
    total = 0.0
    for i in range(len(bars) - 14, len(bars)):
        tr = max(bars[i]["high"] - bars[i]["low"],
                  abs(bars[i]["high"] - bars[i - 1]["close"]),
                  abs(bars[i]["low"] - bars[i - 1]["close"]))
        total += tr
    return total / 14


def atr14_prev(bars):
    if len(bars) < 20:
        return None
    total = 0.0
    for i in range(len(bars) - 19, len(bars) - 5):
        tr = max(bars[i]["high"] - bars[i]["low"],
                  abs(bars[i]["high"] - bars[i - 1]["close"]),
                  abs(bars[i]["low"] - bars[i - 1]["close"]))
        total += tr
    return total / 14


def median(arr):
    if not arr:
        return None
    return statistics.median(arr)


def compute_metrics(raw_bars):
    bars = [{"time": b[0], "open": b[1], "high": b[2], "low": b[3], "close": b[4], "volume": b[5]} for b in raw_bars]
    if len(bars) < 2:
        return None
    last, prev = bars[-1], bars[-2]
    closes = [b["close"] for b in bars]
    volumes = [b["volume"] for b in bars]
    highs = [b["high"] for b in bars]

    last_date = datetime.fromtimestamp(last["time"], tz=timezone.utc)
    one_year_ago = last_date.replace(year=last_date.year - 1)

    hi52 = lo52 = None
    hi, lo = -math.inf, math.inf
    for b in bars:
        d = datetime.fromtimestamp(b["time"], tz=timezone.utc)
        if d >= one_year_ago:
            hi = max(hi, b["high"])
            lo = min(lo, b["low"])
    if hi != -math.inf:
        hi52, lo52 = hi, lo

    v20 = bars[-20:]
    avg_vol = sum(b["volume"] or 0 for b in v20) / len(v20)
    v5, v10 = bars[-5:], bars[-10:]
    avg_vol5 = sum(b["volume"] or 0 for b in v5) / len(v5)
    avg_vol10 = sum(b["volume"] or 0 for b in v10) / len(v10)

    rsi = rsi14(closes)
    ma20, ma50, ma200 = sma(closes, 20), sma(closes, 50), sma(closes, 200)
    atr, atr_prev = atr14(bars), atr14_prev(bars)

    last20 = bars[-20:]
    higher_lows = total_lows = 0
    for i in range(2, len(last20)):
        if last20[i]["low"] > last20[i - 2]["low"]:
            higher_lows += 1
        total_lows += 1

    rsi_prev = rsi14(closes[:-5]) if len(bars) >= 19 else None

    swing_bars = bars[-12:]
    swing_low = min(b["low"] for b in swing_bars) if swing_bars else None

    daily = None
    if prev["close"]:
        daily = (last["close"] - prev["close"]) / prev["close"] * 100

    off_high = None
    if hi52:
        off_high = (last["close"] - hi52) / hi52 * 100

    return {
        "price": last["close"], "daily": daily, "vol": last["volume"], "avgVol": avg_vol,
        "avgVol5": avg_vol5, "avgVol10": avg_vol10, "hi52": hi52, "lo52": lo52, "offHigh": off_high,
        "rsi": rsi, "ma20": ma20, "ma50": ma50, "ma200": ma200, "atr": atr, "atrPrev": atr_prev,
        "higherLows": (higher_lows / total_lows) if total_lows > 0 else None, "rsiPrev": rsi_prev,
        "closes": closes, "volumes": volumes, "highs": highs, "swingLow": swing_low, "bars": bars,
        "asOf": last_date.strftime("%Y-%m-%d"),
    }


def compute_cross_status(raw_bars):
    if not raw_bars or len(raw_bars) < 200 + CROSS_LOOKBACK + 1:
        return None
    closes = [b[4] for b in raw_bars]
    n = len(closes)

    def sma_at(i, length):
        return sum(closes[i - length + 1:i + 1]) / length

    diffs = [sma_at(i, 50) - sma_at(i, 200) for i in range(n - 1 - CROSS_LOOKBACK, n)]
    last = len(diffs) - 1
    ma50, ma200 = sma_at(n - 1, 50), sma_at(n - 1, 200)
    gap_pct = (ma50 - ma200) / ma200 * 100

    for j in range(last, 0, -1):
        if diffs[j] * diffs[j - 1] < 0:
            return {"type": "golden" if diffs[last] > 0 else "death", "status": "crossed",
                    "barsAgo": last - j, "ma50": ma50, "ma200": ma200, "gapPct": gap_pct}

    if abs(gap_pct) <= CROSS_APPROACH_GAP and abs(diffs[last]) < abs(diffs[last - CROSS_CONVERGE_BARS]):
        return {"type": "golden" if diffs[last] < 0 else "death", "status": "approaching",
                "barsAgo": None, "ma50": ma50, "ma200": ma200, "gapPct": gap_pct}
    return None


def compute_rebound_ma50(d):
    price, atr, swing_low = d["price"], d["atr"], d["swingLow"]
    if None in (price, atr, swing_low, d["ma20"], d["ma50"], d["ma200"]):
        return None
    confirmed_uptrend = price > d["ma200"] and d["ma20"] > d["ma50"]
    needs_to_rise = d["ma50"] > price
    if not (confirmed_uptrend and needs_to_rise):
        return None
    reward = d["ma50"] - price
    risk = price - swing_low
    min_risk = MIN_RISK_ATR_MULT * atr
    if risk < min_risk:
        risk = min_risk
    rr = reward / risk if risk > 0 else None
    return {"target": d["ma50"], "reward": reward, "risk": risk, "rr": rr}


def avg_down_day_volume(closes, volumes, start_idx, end_idx):
    total = count = 0
    for i in range(max(start_idx, 1), end_idx):
        if closes[i] < closes[i - 1]:
            total += volumes[i]
            count += 1
    return (total / count) if count > 0 else None


def check_bull_flag(d):
    if d["price"] is None or d["ma50"] is None or d["ma200"] is None:
        return False
    if len(d["volumes"]) < 26:
        return False
    if not (d["price"] > d["ma50"] > d["ma200"]):
        return False
    recent_high = max(d["highs"][-20:])
    if not recent_high:
        return False
    off_high = (recent_high - d["price"]) / recent_high * 100
    if not (BULL_FLAG_MIN_PULLBACK <= off_high <= BULL_FLAG_MAX_PULLBACK):
        return False
    closes, vols, n = d["closes"], d["volumes"], len(d["closes"])
    recent_down_vol = avg_down_day_volume(closes, vols, n - 5, n)
    baseline_down_vol = avg_down_day_volume(closes, vols, n - 25, n - 5)
    if baseline_down_vol is None:
        return False
    sell_vol_ratio = 0 if recent_down_vol is None else recent_down_vol / baseline_down_vol
    if sell_vol_ratio >= BULL_FLAG_MAX_VOL_RATIO:
        return False
    if d["rsi"] is None or not (BULL_FLAG_RSI_MIN <= d["rsi"] <= BULL_FLAG_RSI_MAX):
        return False
    return True


def check_eod_breakout(d):
    # Approximation of "Intraday Breakout Watch" using the most recent completed
    # EOD bar (this tracker only has EOD data.json, not a live intraday feed) --
    # up 5%+ on the day AND >=2x the trailing 20-session MEDIAN volume.
    vols, highs = d["volumes"], d["highs"]
    if len(vols) < 22:
        return False
    today_vol = vols[-1]
    prior20 = vols[-21:-1]
    med_vol = median(prior20)
    if not med_vol:
        return False
    vol_ratio = today_vol / med_vol
    day_pct = d["daily"]
    return day_pct is not None and day_pct >= BREAKOUT_MIN_DAY_PCT and vol_ratio >= BREAKOUT_MIN_VOL_RATIO


def rank_score(items, sym, higher_is_better):
    # items: list of (sym, value)
    match = [v for s, v in items if s == sym]
    if not match:
        return 50.0
    val = match[0]
    sorted_items = sorted(items, key=lambda x: x[1], reverse=higher_is_better)
    idx = next(i for i, (s, _) in enumerate(sorted_items) if s == sym)
    denom = (len(sorted_items) - 1) or 1
    return (1 - idx / denom) * 100


def bounded_score(val, lo, hi):
    if val is None:
        return 50.0
    return max(0, min(100, (val - lo) / (hi - lo) * 100))


def combined_signal(mom, per):
    if mom is None or per is None:
        return "HOLD"
    if mom >= 70 and per >= 70:
        return "STRONG BUY"
    if mom >= 70 and per >= 50:
        return "BUY"
    if mom >= 70 and per < 50:
        return "CAUTION"
    if mom >= 50 and per >= 70:
        return "ACCUMULATE"
    if mom >= 50 and per >= 50:
        return "HOLD"
    if mom >= 50 and per < 50:
        return "WEAKENING"
    if mom < 50 and per >= 70:
        return "WATCH"
    if mom < 50 and per >= 50:
        return "AVOID"
    return "SELL"


def compute_scores(metrics, fund, spy_sym="SPY"):
    spy = metrics.get(spy_sym)
    vals = {k: [] for k in ["ret20", "vsSpy5", "trendAlign", "volTrend", "volCompress", "higherLows", "rsiTraj"]}

    for t, d in metrics.items():
        closes = d["closes"]
        if len(closes) >= 21:
            r20 = (closes[-1] - closes[-21]) / closes[-21] * 100
            vals["ret20"].append((t, r20))
        if spy and len(spy["closes"]) >= 6 and len(closes) >= 6:
            sym5 = (closes[-1] - closes[-6]) / closes[-6] * 100
            spy5 = (spy["closes"][-1] - spy["closes"][-6]) / spy["closes"][-6] * 100
            vals["vsSpy5"].append((t, sym5 - spy5))
        if d["price"] and d["ma20"] and d["ma50"]:
            vals["trendAlign"].append((t, 1 if (d["price"] > d["ma20"] > d["ma50"]) else 0))
        if d["avgVol5"] and d["avgVol10"] and d["avgVol"]:
            vals["volTrend"].append((t, 1 if (d["avgVol5"] > d["avgVol10"] > d["avgVol"]) else 0))
        if d["atr"] is not None and d["atrPrev"] is not None and len(closes) >= 6:
            price_rising = closes[-1] > closes[-6]
            atr_falling = d["atr"] < d["atrPrev"]
            vals["volCompress"].append((t, 1 if (price_rising and atr_falling) else 0))
        if d["higherLows"] is not None:
            vals["higherLows"].append((t, d["higherLows"]))
        if d["rsi"] is not None and d["rsiPrev"] is not None:
            vals["rsiTraj"].append((t, 1 if d["rsi"] > d["rsiPrev"] else 0))

    signals = {}
    for t, d in metrics.items():
        s1 = rank_score(vals["ret20"], t, True)
        s2 = rank_score(vals["vsSpy5"], t, True)
        s3 = bounded_score((d["vol"] / (d["avgVol"] or 1)) if d["avgVol"] else None, 0.5, 3)
        s4 = (100 - abs(d["rsi"] - 60) * 5) if (d["rsi"] is not None and 50 <= d["rsi"] <= 70) else (20 if d["rsi"] is not None else 50)
        s5 = 100 if (d["price"] and d["ma20"] and d["price"] > d["ma20"]) else (0 if (d["price"] and d["ma20"]) else 50)
        s6 = bounded_score((d["price"] - d["lo52"]) / (d["hi52"] - d["lo52"]) * 100, 0, 100) if (d["hi52"] and d["lo52"] and d["hi52"] > d["lo52"]) else 50
        beta = fund.get(t, {}).get("beta")
        s7 = 100 if (beta is not None and beta < 1.2) else (30 if beta is not None else 50)
        mom = s1 * 0.25 + s2 * 0.20 + s3 * 0.15 + s4 * 0.15 + s5 * 0.10 + s6 * 0.10 + s7 * 0.05

        p1 = rank_score(vals["trendAlign"], t, True)
        p2 = rank_score(vals["volTrend"], t, True)
        p3 = rank_score(vals["volCompress"], t, True)
        p4 = rank_score(vals["higherLows"], t, True)
        p5 = rank_score(vals["rsiTraj"], t, True)
        p6 = (80 if (spy and d["daily"] is not None and spy["daily"] is not None and d["daily"] > spy["daily"]) else 30) if spy else 50
        per = p1 * 0.25 + p2 * 0.20 + p3 * 0.15 + p4 * 0.15 + p5 * 0.15 + p6 * 0.10

        signals[t] = {"mom": mom, "per": per, "signal": combined_signal(mom, per)}
    return signals


CATEGORIES = ["Uptrend MA50", "Golden Cross", "Bull Flag Watch", "Breakout Watch (EOD)",
              "Performance: Strong Buy", "Performance: Buy"]


def build_snapshot(raw):
    tickers = raw["tickers"]
    fund = {t: {"beta": v.get("beta")} for t, v in tickers.items()}
    metrics = {}
    for t, v in tickers.items():
        m = compute_metrics(v["bars"])
        if m:
            metrics[t] = m

    scores = compute_scores(metrics, fund)

    membership = {c: [] for c in CATEGORIES}
    for t, d in metrics.items():
        reb = compute_rebound_ma50(d)
        if reb and reb["rr"] is not None and reb["rr"] >= MIN_RR_MA50:
            membership["Uptrend MA50"].append((t, d["price"]))

        cross = compute_cross_status(tickers[t]["bars"])
        if cross and cross["type"] == "golden":
            membership["Golden Cross"].append((t, d["price"]))

        if check_bull_flag(d):
            membership["Bull Flag Watch"].append((t, d["price"]))

        if check_eod_breakout(d):
            membership["Breakout Watch (EOD)"].append((t, d["price"]))

        sig = scores.get(t, {}).get("signal")
        if sig == "STRONG BUY":
            membership["Performance: Strong Buy"].append((t, d["price"]))
        elif sig == "BUY":
            membership["Performance: Buy"].append((t, d["price"]))

    return membership, metrics


def ensure_workbook(path):
    try:
        wb = openpyxl.load_workbook(path)
    except FileNotFoundError:
        wb = Workbook()
        ws = wb.active
        ws.title = "Log"
        ws.append(["Date", "Category", "Ticker", "EntryPrice",
                   "Fwd5D_Ret%", "Fwd5D_Status", "Fwd10D_Ret%", "Fwd10D_Status",
                   "Fwd20D_Ret%", "Fwd20D_Status", "Notes"])
        stats = wb.create_sheet("Stats")
        stats.append(["Category", "N_Logged",
                      "N_5D", "WinRate_5D", "AvgRet_5D",
                      "N_10D", "WinRate_10D", "AvgRet_10D",
                      "N_20D", "WinRate_20D", "AvgRet_20D"])
    if "Log" not in wb.sheetnames:
        ws = wb.create_sheet("Log")
        ws.append(["Date", "Category", "Ticker", "EntryPrice",
                   "Fwd5D_Ret%", "Fwd5D_Status", "Fwd10D_Ret%", "Fwd10D_Status",
                   "Fwd20D_Ret%", "Fwd20D_Status", "Notes"])
    if "Stats" not in wb.sheetnames:
        stats = wb.create_sheet("Stats")
        stats.append(["Category", "N_Logged",
                      "N_5D", "WinRate_5D", "AvgRet_5D",
                      "N_10D", "WinRate_10D", "AvgRet_10D",
                      "N_20D", "WinRate_20D", "AvgRet_20D"])
    return wb


def status(ret):
    if ret is None:
        return None
    if ret > UP_THRESH:
        return "UP"
    if ret < DOWN_THRESH:
        return "DOWN"
    return "FLAT"


def resolve_open_rows(wb, metrics_by_ticker_bars, today_str):
    ws = wb["Log"]
    headers = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(headers)}
    resolved_counts = {5: 0, 10: 0, 20: 0}

    for row in ws.iter_rows(min_row=2):
        date_cell = row[col["Date"] - 1]
        ticker_cell = row[col["Ticker"] - 1]
        if date_cell.value is None or date_cell.value == today_str:
            continue
        ticker = ticker_cell.value
        bars = metrics_by_ticker_bars.get(ticker)
        if not bars:
            continue
        # find bar index matching the log date (or first bar after it)
        log_date = str(date_cell.value)[:10]
        idx = None
        for i, b in enumerate(bars):
            bd = datetime.fromtimestamp(b[0], tz=timezone.utc).strftime("%Y-%m-%d")
            if bd >= log_date:
                idx = i
                break
        if idx is None:
            continue
        entry_close = bars[idx][4]
        for h in HORIZONS:
            ret_col = col[f"Fwd{h}D_Ret%"]
            stat_col = col[f"Fwd{h}D_Status"]
            if row[ret_col - 1].value is not None:
                continue  # already resolved
            target_idx = idx + h
            if target_idx < len(bars):
                fwd_close = bars[target_idx][4]
                ret = (fwd_close - entry_close) / entry_close * 100
                row[ret_col - 1].value = round(ret, 2)
                row[stat_col - 1].value = status(ret)
                resolved_counts[h] += 1
    return resolved_counts


def append_new_rows(wb, membership, today_str):
    ws = wb["Log"]
    existing = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        existing.add((str(row[0])[:10], row[1], row[2]))
    added = 0
    for cat, items in membership.items():
        for ticker, price in items:
            key = (today_str, cat, ticker)
            if key in existing:
                continue
            ws.append([today_str, cat, ticker, round(price, 2), None, None, None, None, None, None, ""])
            added += 1
    return added


# ---------------------------------------------------------------------------
# v2 additions (2026-09-23)
#  - Drops an incomplete (intraday) last bar so signals are only ever computed on
#    completed daily bars. data.json is refreshed intraday; before this fix a run
#    during US hours logged "EOD" signals from a partial bar.
#  - Backfills missing trading days since the first logged date (Notes=BACKFILL).
#  - Flags existing rows whose EntryPrice doesn't match that day's final close
#    (Notes=PARTIAL-BAR) and excludes them from Stats.
#  - Adds benchmark-relative and next-open-entry columns (additive; the original
#    close-to-close Fwd columns and UP/DOWN/FLAT thresholds are unchanged).
# ---------------------------------------------------------------------------
from datetime import timedelta

COMPLETE_AFTER_UTC_HOUR = 21.25   # bar D is complete once asof >= D 21:15 UTC (covers EDT and EST closes)
SLIPPAGE_BPS_PER_SIDE = 5          # haircut for the next-open entry variant
BENCH = "SPY"
BACKFILL_START = "2026-09-10"   # first day of the original forward test
PARTIAL_TOL_PCT = 0.05
EXTRA_COLS = []
for _h in HORIZONS:
    EXTRA_COLS += [f"Fwd{_h}D_vsSPY%", f"Fwd{_h}D_NextOpenNet%"]


def bar_date(b):
    # bars are stamped at 04:00 UTC = midnight US/Eastern of the session date
    return datetime.fromtimestamp(b[0], tz=timezone.utc).strftime("%Y-%m-%d")


def drop_incomplete_bars(raw):
    asof = datetime.fromisoformat(raw["asof"].replace("Z", "+00:00"))
    dropped = 0
    for v in raw["tickers"].values():
        bars = v.get("bars") or []
        if not bars:
            continue
        d = datetime.strptime(bar_date(bars[-1]), "%Y-%m-%d").replace(tzinfo=timezone.utc)
        if asof < d + timedelta(hours=COMPLETE_AFTER_UTC_HOUR):
            v["bars"] = bars[:-1]
            dropped += 1
    return dropped


def truncate_raw(raw, date_str):
    return {"asof": raw["asof"], "tickers": {
        t: {**v, "bars": [b for b in v["bars"] if bar_date(b) <= date_str]}
        for t, v in raw["tickers"].items()}}


def ensure_extra_cols(ws):
    headers = [c.value for c in ws[1]]
    for h in EXTRA_COLS:
        if h not in headers:
            ws.cell(row=1, column=len(headers) + 1, value=h)
            headers.append(h)
    return {h: i + 1 for i, h in enumerate(headers)}


def idx_for_date(bars, d):
    for i, b in enumerate(bars):
        if bar_date(b) == d:
            return i
    return None


def flag_and_extend(ws, bars_by_ticker, col):
    """Flag partial-bar rows and fill the additive benchmark / next-open columns."""
    flagged = 0
    spy = bars_by_ticker.get(BENCH) or []
    spy_idx = {bar_date(b): i for i, b in enumerate(spy)}
    for row in ws.iter_rows(min_row=2):
        d = row[col["Date"] - 1].value
        t = row[col["Ticker"] - 1].value
        if d is None:
            continue
        d = str(d)[:10]
        bars = bars_by_ticker.get(t)
        if not bars:
            continue
        i = idx_for_date(bars, d)
        if i is None:
            continue
        close_d = bars[i][4]
        notes = row[col["Notes"] - 1]
        ep = row[col["EntryPrice"] - 1].value
        if ep and not notes.value and abs(ep - close_d) > max(close_d * PARTIAL_TOL_PCT / 100, 0.006):
            notes.value = "PARTIAL-BAR"
            flagged += 1
        si = spy_idx.get(d)
        for h in HORIZONS:
            vs_c, no_c = col[f"Fwd{h}D_vsSPY%"], col[f"Fwd{h}D_NextOpenNet%"]
            if i + h < len(bars):
                fwd = bars[i + h][4]
                if row[vs_c - 1].value is None and si is not None and si + h < len(spy):
                    r_t = (fwd - close_d) / close_d * 100
                    r_s = (spy[si + h][4] - spy[si][4]) / spy[si][4] * 100
                    row[vs_c - 1].value = round(r_t - r_s, 2)
                if row[no_c - 1].value is None:
                    o = bars[i + 1][1]
                    cost = 2 * SLIPPAGE_BPS_PER_SIDE / 100
                    row[no_c - 1].value = round((fwd - o) / o * 100 - cost, 2)
    return flagged


def update_stats(wb):
    log_ws = wb["Log"]
    if "Stats" in wb.sheetnames:
        del wb["Stats"]
    st = wb.create_sheet("Stats")
    headers = [c.value for c in log_ws[1]]
    col = {h: i for i, h in enumerate(headers)}
    rows = [r for r in log_ws.iter_rows(min_row=2, values_only=True)
            if r[0] is not None and (r[col["Notes"]] or "") != "PARTIAL-BAR"]
    hdr = ["Category", "N_Logged"]
    for h in HORIZONS:
        hdr += [f"N_{h}D", f"WinRate_{h}D", f"AvgRet_{h}D", f"AvgVsSPY_{h}D",
                f"BeatSPY%_{h}D", f"AvgNextOpenNet_{h}D"]
    st.append(hdr)

    def avg(xs):
        return round(sum(xs) / len(xs), 2) if xs else None

    for cat in CATEGORIES:
        cr = [r for r in rows if r[col["Category"]] == cat]
        out = [cat, len(cr)]
        for h in HORIZONS:
            rets = [r[col[f"Fwd{h}D_Ret%"]] for r in cr if r[col[f"Fwd{h}D_Ret%"]] is not None]
            vs = [r[col[f"Fwd{h}D_vsSPY%"]] for r in cr if r[col[f"Fwd{h}D_vsSPY%"]] is not None]
            no = [r[col[f"Fwd{h}D_NextOpenNet%"]] for r in cr if r[col[f"Fwd{h}D_NextOpenNet%"]] is not None]
            n = len(rets)
            out += [n,
                    round(sum(1 for x in rets if x > 0) / n * 100, 1) if n else None,
                    avg(rets), avg(vs),
                    round(sum(1 for x in vs if x > 0) / len(vs) * 100, 1) if vs else None,
                    avg(no)]
        st.append(out)
    st.append([])
    st.append(["Notes: rows flagged PARTIAL-BAR (logged from an intraday bar) are excluded. "
               "Categories with N < 20 at a horizon are insufficient sample. "
               "NextOpenNet = entry at next session open, exit at close h sessions after the "
               f"signal day, minus {2*SLIPPAGE_BPS_PER_SIDE} bps round-trip."])


def main():
    raw = fetch_json(DATA_URL)
    asof = raw.get("asof", "")
    dropped = drop_incomplete_bars(raw)
    bars_by_ticker = {t: v["bars"] for t, v in raw["tickers"].items()}
    spy_dates = [bar_date(b) for b in bars_by_ticker.get(BENCH, [])]
    last_complete = spy_dates[-1]

    wb = ensure_workbook(TRACKER_PATH)
    ws = wb["Log"]
    col = ensure_extra_cols(ws)
    rebuild = ws.max_row <= 1
    logged_dates = sorted({str(r[0])[:10] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]})
    first = min(logged_dates[:1] + [BACKFILL_START])
    todo = [d for d in spy_dates if first <= d <= last_complete and d not in logged_dates]

    added_by_date = {}
    for d in todo:
        membership, _ = build_snapshot(truncate_raw(raw, d))
        n0 = ws.max_row
        append_new_rows(wb, membership, d)
        if d != last_complete and not rebuild:
            for r in range(n0 + 1, ws.max_row + 1):
                ws.cell(row=r, column=col["Notes"], value="BACKFILL")
        added_by_date[d] = ws.max_row - n0

    resolved = resolve_open_rows(wb, bars_by_ticker, last_complete)
    flagged = flag_and_extend(ws, bars_by_ticker, col)
    update_stats(wb)
    wb.save(TRACKER_PATH)

    print(f"asof={asof}  last_complete_bar={last_complete}  incomplete_bars_dropped={dropped}")
    print(f"rows_added_by_date={added_by_date}")
    print(f"resolved={resolved}  newly_flagged_partial={flagged}")
    if last_complete in todo:
        membership, _ = build_snapshot(truncate_raw(raw, last_complete))
        for cat in CATEGORIES:
            print(f"{cat}: {[t for t, _ in membership[cat]]}")
    if EXPORT_DIR:
        export_csvs(wb, EXPORT_DIR, last_complete, asof)
    if last_complete not in todo:
        print("No new completed session to log (data.json not yet refreshed past last logged date).")


def export_csvs(wb, out_dir, last_complete, asof):
    import csv, os
    os.makedirs(out_dir, exist_ok=True)
    rows = list(wb["Log"].iter_rows(values_only=True))
    with open(os.path.join(out_dir, "log.csv"), "w", newline="") as f:
        csv.writer(f).writerows(rows)
    with open(os.path.join(out_dir, "stats.csv"), "w", newline="") as f:
        w = csv.writer(f)
        w.writerow([f"BCC forward-test stats | data asof {asof} | last completed session {last_complete}"])
        w.writerows(r for r in wb["Stats"].iter_rows(values_only=True) if any(v is not None for v in r))
    with open(os.path.join(out_dir, f"membership_{last_complete}.csv"), "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(rows[0][:4])
        w.writerows(r[:4] for r in rows[1:] if str(r[0])[:10] == last_complete)
    print(f"exported to {out_dir}")


if __name__ == "__main__":
    main()
